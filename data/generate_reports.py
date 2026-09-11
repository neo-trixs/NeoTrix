#!/usr/bin/env python3
"""Generate Excel (4 sheets) and PDF reports from contracts_full.json."""
import json
import os
from collections import defaultdict
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter

# ─── Load data ───
DATA_PATH = "data/parsed/contracts_full.json"
OUTPUT_DIR = "output"
os.makedirs(OUTPUT_DIR, exist_ok=True)

with open(DATA_PATH, encoding="utf-8") as f:
    contracts = json.load(f)

# ─── Flatten products ───
all_products = []
for c in contracts:
    for p in c["products"]:
        all_products.append({
            "contract_no": c["contract_no"],
            "buyer_name": c["buyer_name"],
            "seller_name": c["seller_name"],
            "salesperson": c["salesperson"],
            "contract_date": c["contract_date"],
            **p
        })

# ─── Helper: style header ───
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(name="微软雅黑", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
ALT_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
TOTAL_FONT = Font(name="微软雅黑", bold=True, size=11)
TOTAL_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
MONEY_FMT = '#,##0.00'
INT_FMT = '#,##0'

def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER

def style_data_row(ws, row, ncols, alt=False):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")
        if alt:
            cell.fill = ALT_FILL

def auto_width(ws, ncols, max_width=30):
    for col in range(1, ncols + 1):
        max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 4, max_width)

# ═══════════════════════════════════════════════════════════
# EXCEL REPORT
# ═══════════════════════════════════════════════════════════
wb = Workbook()

# ─── Sheet1: 合同明细 ───
ws1 = wb.active
ws1.title = "合同明细"
headers1 = ["合同编号", "买方", "卖方", "业务员", "合同日期", "总金额", "总数量", "付款条款", "执行标准", "产品数"]
ws1.append(headers1)
style_header(ws1, 1, len(headers1))
for i, c in enumerate(contracts, start=2):
    ws1.append([
        c["contract_no"], c["buyer_name"], c["seller_name"], c["salesperson"],
        c["contract_date"], c["total_amount"], c["total_qty"],
        c["payment_terms"], c["execution_standard"], len(c["products"])
    ])
    style_data_row(ws1, i, len(headers1), alt=(i % 2 == 0))
    ws1.cell(row=i, column=6).number_format = MONEY_FMT
    ws1.cell(row=i, column=7).number_format = INT_FMT

# Total row
total_row = len(contracts) + 2
ws1.cell(row=total_row, column=1, value="合计").font = TOTAL_FONT
ws1.cell(row=total_row, column=1).fill = TOTAL_FILL
ws1.cell(row=total_row, column=1).border = THIN_BORDER
ws1.cell(row=total_row, column=6, value=f"=SUM(F2:F{total_row-1})")
ws1.cell(row=total_row, column=6).font = TOTAL_FONT
ws1.cell(row=total_row, column=6).fill = TOTAL_FILL
ws1.cell(row=total_row, column=6).number_format = MONEY_FMT
ws1.cell(row=total_row, column=6).border = THIN_BORDER
ws1.cell(row=total_row, column=10, value=f"=SUM(J2:J{total_row-1})")
ws1.cell(row=total_row, column=10).font = TOTAL_FONT
ws1.cell(row=total_row, column=10).fill = TOTAL_FILL
ws1.cell(row=total_row, column=10).border = THIN_BORDER
auto_width(ws1, len(headers1))

# ─── Sheet2: 产品明细 ───
ws2 = wb.create_sheet("产品明细")
headers2 = ["合同编号", "买方", "卖方", "业务员", "合同日期", "产品名称", "直径", "规格", "数量", "单位", "单价", "金额", "单重", "总重"]
ws2.append(headers2)
style_header(ws2, 1, len(headers2))
for i, p in enumerate(all_products, start=2):
    ws2.append([
        p["contract_no"], p["buyer_name"], p["seller_name"], p["salesperson"],
        p["contract_date"], p["name"], p["diameter"], p["config"],
        p["qty"], p["unit"], p["price"], p["amount"], p["unit_weight"], p["total_weight"]
    ])
    style_data_row(ws2, i, len(headers2), alt=(i % 2 == 0))
    ws2.cell(row=i, column=11).number_format = MONEY_FMT
    ws2.cell(row=i, column=12).number_format = MONEY_FMT
    ws2.cell(row=i, column=13).number_format = MONEY_FMT
    ws2.cell(row=i, column=14).number_format = MONEY_FMT
auto_width(ws2, len(headers2))

# ─── Sheet3: 业务员汇总 ───
ws3 = wb.create_sheet("业务员汇总")
# Build stats including zero-contract salespeople
ALL_SALESPEOPLE = sorted(set(c["salesperson"] for c in contracts))
ZERO_CONTRACT = ["韩冰", "谢瑶"]
for z in ZERO_CONTRACT:
    if z not in ALL_SALESPEOPLE:
        ALL_SALESPEOPLE.append(z)
ALL_SALESPEOPLE = sorted(ALL_SALESPEOPLE)

sp_stats = defaultdict(lambda: {"contracts": 0, "amount": 0, "products": 0, "sellers": set()})
for c in contracts:
    sp = c["salesperson"]
    sp_stats[sp]["contracts"] += 1
    sp_stats[sp]["amount"] += c["total_amount"]
    sp_stats[sp]["products"] += len(c["products"])
    sp_stats[sp]["sellers"].add(c["seller_name"])

headers3 = ["业务员", "合同数", "总金额", "产品数", "供应商数"]
ws3.append(headers3)
style_header(ws3, 1, len(headers3))
for i, sp in enumerate(ALL_SALESPEOPLE, start=2):
    s = sp_stats[sp]
    ws3.append([sp, s["contracts"], round(s["amount"], 2), s["products"], len(s["sellers"])])
    style_data_row(ws3, i, len(headers3), alt=(i % 2 == 0))
    ws3.cell(row=i, column=3).number_format = MONEY_FMT
    ws3.cell(row=i, column=2).number_format = INT_FMT
    ws3.cell(row=i, column=4).number_format = INT_FMT
    ws3.cell(row=i, column=5).number_format = INT_FMT
auto_width(ws3, len(headers3))

# ─── Sheet4: 供应商汇总 ───
ws4 = wb.create_sheet("供应商汇总")
seller_stats = defaultdict(lambda: {"contracts": 0, "amount": 0})
for c in contracts:
    seller_stats[c["seller_name"]]["contracts"] += 1
    seller_stats[c["seller_name"]]["amount"] += c["total_amount"]

headers4 = ["供应商", "合同数", "总金额"]
ws4.append(headers4)
style_header(ws4, 1, len(headers4))
for i, (seller, s) in enumerate(sorted(seller_stats.items(), key=lambda x: -x[1]["amount"]), start=2):
    ws4.append([seller, s["contracts"], round(s["amount"], 2)])
    style_data_row(ws4, i, len(headers4), alt=(i % 2 == 0))
    ws4.cell(row=i, column=3).number_format = MONEY_FMT
auto_width(ws4, len(headers4))

EXCEL_PATH = os.path.join(OUTPUT_DIR, "2026年合同数据汇总.xlsx")
wb.save(EXCEL_PATH)
print(f"Excel saved: {EXCEL_PATH}")

# ═══════════════════════════════════════════════════════════
# PDF REPORT
# ═══════════════════════════════════════════════════════════
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
import tempfile

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm, cm
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    PageBreak, Image
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

# Register CID font for Chinese
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

# Colors
PRIMARY = HexColor("#1F4E79")
ACCENT = HexColor("#4472C4")
LIGHT_BG = HexColor("#D9E2F3")
WHITE = HexColor("#FFFFFF")
BLACK = HexColor("#000000")
GRAY = HexColor("#666666")

# Styles
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="CNTitle", fontName="STSong-Light", fontSize=28, leading=36, alignment=1, textColor=PRIMARY, spaceAfter=6))
styles.add(ParagraphStyle(name="CNSubtitle", fontName="STSong-Light", fontSize=14, leading=20, alignment=1, textColor=ACCENT, spaceAfter=20))
styles.add(ParagraphStyle(name="CNHeading", fontName="STSong-Light", fontSize=16, leading=22, textColor=PRIMARY, spaceBefore=16, spaceAfter=10))
styles.add(ParagraphStyle(name="CNSmall", fontName="STSong-Light", fontSize=9, leading=13, textColor=GRAY))
styles.add(ParagraphStyle(name="CNBody", fontName="STSong-Light", fontSize=10, leading=14))
styles.add(ParagraphStyle(name="CNCell", fontName="STSong-Light", fontSize=8, leading=11))

PDF_PATH = os.path.join(OUTPUT_DIR, "2026年合同数据报告.pdf")
doc = SimpleDocTemplate(PDF_PATH, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm, leftMargin=1.5*cm, rightMargin=1.5*cm)
story = []

# ─── 1. Cover ───
total_amount = sum(c["total_amount"] for c in contracts)
story.append(Spacer(1, 80))
story.append(Paragraph("2026年合同数据报告", styles["CNTitle"]))
story.append(Spacer(1, 10))
story.append(Paragraph("Contract Data Annual Report", styles["CNSubtitle"]))
story.append(Spacer(1, 30))
cover_data = [
    ["报告日期", datetime.now().strftime("%Y-%m-%d")],
    ["合同总数", f"{len(contracts)} 份"],
    ["产品总数", f"{len(all_products)} 条"],
    ["业务员数", f"{len(ALL_SALESPEOPLE)} 人"],
    ["总金额", f"¥{total_amount:,.2f}"],
]
cover_table = Table(cover_data, colWidths=[120, 200])
cover_table.setStyle(TableStyle([
    ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
    ("FONTSIZE", (0, 0), (-1, -1), 12),
    ("TEXTCOLOR", (0, 0), (0, -1), PRIMARY),
    ("TEXTCOLOR", (1, 0), (1, -1), BLACK),
    ("ALIGN", (0, 0), (0, -1), "RIGHT"),
    ("ALIGN", (1, 0), (1, -1), "LEFT"),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ("TOPPADDING", (0, 0), (-1, -1), 8),
    ("LINEBELOW", (0, 0), (-1, -1), 0.5, GRAY),
]))
story.append(cover_table)
story.append(PageBreak())

# ─── 2. Salesperson Summary ───
story.append(Paragraph("一、业务员汇总", styles["CNHeading"]))
story.append(Paragraph(f"共 {len(ALL_SALESPEOPLE)} 名业务员，其中 {len(ZERO_CONTRACT)} 人无合同记录", styles["CNSmall"]))
story.append(Spacer(1, 8))

sp_data = [["序号", "业务员", "合同数", "总金额", "产品数", "供应商数"]]
for idx, sp in enumerate(ALL_SALESPEOPLE, start=1):
    s = sp_stats[sp]
    sp_data.append([str(idx), sp, str(s["contracts"]), f"¥{s['amount']:,.2f}", str(s["products"]), str(len(s["sellers"]))])

sp_table = Table(sp_data, colWidths=[35, 70, 55, 120, 55, 55])
sp_style = [
    ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
    ("FONTSIZE", (0, 0), (-1, -1), 8),
    ("BACKGROUND", (0, 0), (-1, 0), PRIMARY),
    ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
    ("ALIGN", (0, 0), (-1, 0), "CENTER"),
    ("ALIGN", (2, 1), (-1, -1), "CENTER"),
    ("ALIGN", (3, 1), (3, -1), "RIGHT"),
    ("GRID", (0, 0), (-1, -1), 0.5, GRAY),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT_BG]),
    ("TOPPADDING", (0, 0), (-1, -1), 4),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
]
# Highlight zero-contract rows
for idx, sp in enumerate(ALL_SALESPEOPLE, start=1):
    if sp in ZERO_CONTRACT:
        sp_style.append(("TEXTCOLOR", (0, idx), (-1, idx), HexColor("#C00000")))
        sp_style.append(("FONTSIZE", (0, idx), (-1, idx), 7))
sp_table.setStyle(TableStyle(sp_style))
story.append(sp_table)
story.append(PageBreak())

# ─── 3. Product Type Analysis ───
story.append(Paragraph("二、产品类型分析", styles["CNHeading"]))
story.append(Spacer(1, 8))

product_stats = defaultdict(lambda: {"qty": 0, "amount": 0, "count": 0})
for p in all_products:
    product_stats[p["name"]]["qty"] += p["qty"]
    product_stats[p["name"]]["amount"] += p["amount"]
    product_stats[p["name"]]["count"] += 1

sorted_products = sorted(product_stats.items(), key=lambda x: -x[1]["amount"])

prod_data = [["产品名称", "产品条数", "总数量", "总金额"]]
for name, s in sorted_products:
    prod_data.append([name, str(s["count"]), str(s["qty"]), f"¥{s['amount']:,.2f}"])
prod_table = Table(prod_data, colWidths=[100, 70, 70, 120])
prod_table.setStyle(TableStyle([
    ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
    ("FONTSIZE", (0, 0), (-1, -1), 8),
    ("BACKGROUND", (0, 0), (-1, 0), PRIMARY),
    ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
    ("ALIGN", (1, 0), (-1, -1), "CENTER"),
    ("ALIGN", (3, 1), (3, -1), "RIGHT"),
    ("GRID", (0, 0), (-1, -1), 0.5, GRAY),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT_BG]),
    ("TOPPADDING", (0, 0), (-1, -1), 4),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
]))
story.append(prod_table)
story.append(Spacer(1, 12))

# Product chart
tmp_chart = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
fig, ax = plt.subplots(figsize=(8, 4))
chart_names = [n for n, _ in sorted_products[:10]]
chart_amounts = [s["amount"] for _, s in sorted_products[:10]]
bars = ax.barh(chart_names[::-1], chart_amounts[::-1], color="#4472C4", edgecolor="#1F4E79")
ax.set_xlabel("Total Amount", fontsize=10)
ax.set_title("Top 10 Product Types by Amount", fontsize=12, color="#1F4E79")
ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: f"{x/1e6:.1f}M"))
plt.tight_layout()
plt.savefig(tmp_chart.name, dpi=150, bbox_inches="tight")
plt.close()
story.append(Image(tmp_chart.name, width=480, height=240))
story.append(PageBreak())

# ─── 4. Supplier TOP15 ───
story.append(Paragraph("三、供应商 TOP15", styles["CNHeading"]))
story.append(Spacer(1, 8))

sorted_sellers = sorted(seller_stats.items(), key=lambda x: -x[1]["amount"])[:15]
seller_data = [["排名", "供应商", "合同数", "总金额"]]
for rank, (seller, s) in enumerate(sorted_sellers, start=1):
    seller_data.append([str(rank), seller, str(s["contracts"]), f"¥{s['amount']:,.2f}"])
seller_table = Table(seller_data, colWidths=[35, 180, 55, 120])
seller_table.setStyle(TableStyle([
    ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
    ("FONTSIZE", (0, 0), (-1, -1), 9),
    ("BACKGROUND", (0, 0), (-1, 0), PRIMARY),
    ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
    ("ALIGN", (0, 0), (0, -1), "CENTER"),
    ("ALIGN", (2, 0), (-1, -1), "CENTER"),
    ("ALIGN", (3, 1), (3, -1), "RIGHT"),
    ("GRID", (0, 0), (-1, -1), 0.5, GRAY),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT_BG]),
    ("TOPPADDING", (0, 0), (-1, -1), 5),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
]))
story.append(seller_table)
story.append(Spacer(1, 12))

# Supplier chart
tmp_seller = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
fig2, ax2 = plt.subplots(figsize=(8, 4))
s_names = [s for s, _ in sorted_sellers]
s_amounts = [v["amount"] for _, v in sorted_sellers]
ax2.barh(s_names[::-1], [a/1e6 for a in s_amounts[::-1]], color="#70AD47", edgecolor="#375623")
ax2.set_xlabel("Total Amount (Million)", fontsize=10)
ax2.set_title("Top 15 Suppliers by Amount", fontsize=12, color="#1F4E79")
plt.tight_layout()
plt.savefig(tmp_seller.name, dpi=150, bbox_inches="tight")
plt.close()
story.append(Image(tmp_seller.name, width=480, height=240))
story.append(PageBreak())

# ─── 5. Monthly Trend ───
story.append(Paragraph("四、月度趋势分析", styles["CNHeading"]))
story.append(Spacer(1, 8))

monthly = defaultdict(lambda: {"contracts": 0, "amount": 0, "products": 0})
for c in contracts:
    month = c["contract_date"][:7]  # YYYY-MM
    monthly[month]["contracts"] += 1
    monthly[month]["amount"] += c["total_amount"]
    monthly[month]["products"] += len(c["products"])

sorted_months = sorted(monthly.items())
month_data = [["月份", "合同数", "产品数", "总金额"]]
for m, s in sorted_months:
    month_data.append([m, str(s["contracts"]), str(s["products"]), f"¥{s['amount']:,.2f}"])
month_table = Table(month_data, colWidths=[80, 60, 60, 120])
month_table.setStyle(TableStyle([
    ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
    ("FONTSIZE", (0, 0), (-1, -1), 9),
    ("BACKGROUND", (0, 0), (-1, 0), PRIMARY),
    ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ("ALIGN", (3, 1), (3, -1), "RIGHT"),
    ("GRID", (0, 0), (-1, -1), 0.5, GRAY),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT_BG]),
    ("TOPPADDING", (0, 0), (-1, -1), 5),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
]))
story.append(month_table)
story.append(Spacer(1, 12))

# Monthly chart
tmp_monthly = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
fig3, ax3 = plt.subplots(figsize=(8, 4))
m_labels = [m for m, _ in sorted_months]
m_amounts = [s["amount"]/1e6 for _, s in sorted_months]
m_contracts = [s["contracts"] for _, s in sorted_months]
ax3_twin = ax3.twinx()
ax3.bar(m_labels, m_amounts, color="#4472C4", alpha=0.7, label="Amount (M)")
ax3_twin.plot(m_labels, m_contracts, color="#C00000", marker="o", linewidth=2, label="Contracts")
ax3.set_xlabel("Month", fontsize=10)
ax3.set_ylabel("Amount (Million)", fontsize=10, color="#4472C4")
ax3_twin.set_ylabel("Contract Count", fontsize=10, color="#C00000")
ax3.set_title("Monthly Contract Trend", fontsize=12, color="#1F4E79")
ax3.set_xticks(range(len(m_labels)))
ax3.set_xticklabels(m_labels, rotation=45, fontsize=8)
lines1, labels1 = ax3.get_legend_handles_labels()
lines2, labels2 = ax3_twin.get_legend_handles_labels()
ax3.legend(lines1 + lines2, labels1 + labels2, loc="upper left", fontsize=8)
plt.tight_layout()
plt.savefig(tmp_monthly.name, dpi=150, bbox_inches="tight")
plt.close()
story.append(Image(tmp_monthly.name, width=480, height=240))

# Build PDF
doc.build(story)
print(f"PDF saved: {PDF_PATH}")

# Cleanup temp files
for tmp in [tmp_chart, tmp_seller, tmp_monthly]:
    try:
        os.unlink(tmp.name)
    except:
        pass

print("Done!")
