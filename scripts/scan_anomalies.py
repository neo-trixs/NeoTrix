#!/usr/bin/env python3
"""
Step 1: Quick scan to find anomalies
"""
import os, json, re
from pathlib import Path
from openpyxl import load_workbook

CONTRACTS_DIR = Path("/tmp/smb_contracts")

def quick_parse(filepath):
    """Minimal parse - just get total and product sum."""
    try:
        wb = load_workbook(filepath, data_only=True, read_only=True)
    except:
        return None

    # Find sheet
    sn = None
    for n in wb.sheetnames:
        if '合同' in n:
            sn = n
            break
    if not sn:
        sn = wb.sheetnames[0] if wb.sheetnames else None
    if not sn:
        wb.close()
        return None

    ws = wb[sn]
    total = None
    prod_sum = 0.0
    prod_count = 0

    # Scan rows
    for row in ws.iter_rows(min_row=1, max_row=25, max_col=10):
        a = row[0].value if row[0].value else ''
        # Find 合计
        if '合计' in str(a):
            for cell in row[1:]:
                if cell.value is not None:
                    try:
                        total = float(cell.value)
                        break
                    except:
                        pass
        # Find products (amounts in col G/H)
        elif prod_count < 20:  # limit
            for ci in [6, 7, 5]:  # 0-indexed: G=6, H=7, F=5
                if ci < len(row) and row[ci].value is not None:
                    try:
                        v = float(row[ci].value)
                        if v > 0:
                            prod_sum += v
                            prod_count += 1
                            break
                    except:
                        pass

    wb.close()
    if total is not None and prod_sum > 0:
        return {"total": total, "prod_sum": prod_sum}
    return None

# Scan all files
anomalies = []
all_files = []
for person in sorted(CONTRACTS_DIR.iterdir()):
    if not person.is_dir(): continue
    for cd in sorted(person.iterdir()):
        if not cd.is_dir(): continue
        for f in cd.iterdir():
            if f.suffix.lower() in ('.xlsx','.xls') and not f.name.startswith('~$'):
                all_files.append((person.name, cd.name, f))

print(f"Scanning {len(all_files)} files...")
for i, (p, d, f) in enumerate(all_files):
    if i % 50 == 0:
        print(f"  {i}/{len(all_files)}...")
    r = quick_parse(f)
    if r and r["prod_sum"] > r["total"]:
        anomalies.append({"person": p, "dir": d, "file": f.name, "path": str(f),
                          "total": r["total"], "prod_sum": r["prod_sum"]})

print(f"\nFound {len(anomalies)} anomalies:")
for a in anomalies:
    print(f"  {a['person']}/{a['file']}: total={a['total']}, prod_sum={a['prod_sum']:.2f}")

# Save anomalies for step 2
with open("/tmp/anomalies.json", 'w') as f:
    json.dump(anomalies, f, ensure_ascii=False, indent=2)
print(f"\nSaved to /tmp/anomalies.json")
